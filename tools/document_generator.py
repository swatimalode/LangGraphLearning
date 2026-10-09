from pathlib import Path
from typing import Any
import re
import uuid

from docx import Document
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer


OUTPUT_DIR = Path("generated_documents")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def _safe_filename(filename: str) -> str:
    """Remove unsafe path characters and keep only a simple filename."""
    filename = Path(filename).name
    filename = re.sub(r"[^A-Za-z0-9._ -]", "", filename).strip()
    filename = filename.replace(" ", "_")
    return filename or "generated_document"


def generate_document(
    filename: str,
    file_type: str,
    content: str = "",
    table_data: list[list[Any]] | None = None,
) -> dict[str, str]:
    """
    Create a PDF, DOCX, TXT, or XLSX file.

    Args:
        filename: Desired file name, with or without extension.
        file_type: pdf, docx, txt, or xlsx.
        content: Text content for PDF/DOCX/TXT; also used as an optional
                 title/introduction for XLSX.
        table_data: Optional spreadsheet rows, e.g.
                    [["Name", "Status"], ["Report A", "Complete"]].

    Returns:
        Dictionary containing status, filename, and relative file path.
    """
    normalized_type = file_type.lower().strip().lstrip(".")
    supported_types = {"pdf", "docx", "txt", "xlsx"}
    if normalized_type not in supported_types:
        raise ValueError("Unsupported file type. Use pdf, docx, txt, or xlsx.")

    safe_name = _safe_filename(filename)
    # Replace a mismatched supplied extension with the requested format.
    stem = Path(safe_name).stem or "generated_document"
    output_path = OUTPUT_DIR / f"{stem}_{uuid.uuid4().hex[:8]}.{normalized_type}"

    if normalized_type == "txt":
        output_path.write_text(content, encoding="utf-8")

    elif normalized_type == "docx":
        document = Document()
        for line in content.splitlines() or [""]:
            if line.strip():
                document.add_paragraph(line)
            else:
                document.add_paragraph("")
        document.save(output_path)

    elif normalized_type == "pdf":
        styles = getSampleStyleSheet()
        story = []
        for line in content.splitlines() or [""]:
            if not line.strip():
                story.append(Spacer(1, 8))
                continue
            # Escape XML characters so text such as "A & B" works in a PDF.
            escaped = (
                line.replace("&", "&amp;")
                    .replace("<", "&lt;")
                    .replace(">", "&gt;")
            )
            story.append(Paragraph(escaped, styles["BodyText"]))
            story.append(Spacer(1, 6))
        pdf = SimpleDocTemplate(str(output_path), pagesize=A4)
        pdf.build(story)

    elif normalized_type == "xlsx":
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Generated Data"

        rows = table_data or []
        if content.strip():
            worksheet.append([content.strip()])
            worksheet.append([])
        for row in rows:
            worksheet.append([str(value) if value is not None else "" for value in row])

        if rows:
            # Treat the first row as headers for basic formatting.
            for cell in worksheet[worksheet.max_row - len(rows) + 1]:
                cell.font = cell.font.copy(bold=True, color="FFFFFF")
                cell.fill = __import__("openpyxl").styles.PatternFill(
                    fill_type="solid", fgColor="1F4E78"
                )
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions

        # Give columns a sensible width, capped to avoid enormous columns.
        for column_cells in worksheet.columns:
            column_letter = column_cells[0].column_letter
            max_length = max(
                (len(str(cell.value)) for cell in column_cells if cell.value is not None),
                default=10,
            )
            worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 40)

        workbook.save(output_path)

    return {
        "status": "success",
        "filename": output_path.name,
        "file_path": str(output_path),
        "file_type": normalized_type,
    }
